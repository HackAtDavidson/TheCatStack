from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

from .core import eligible, merge, posted_recently, read_json, undergraduate_internship, write_json
from .catalog import write_catalog
from .digest import prepare_digest
from .sources import collect_source
from .screening import screen_records
from .workbook import export_csv, export_workbook, read_review


def collect(config: dict, config_root: Path, state_path: Path, output: Path, today,
            automate: bool = False, refresh_pages: bool = False) -> int:
    sources = [source for source in config["sources"] if source.get("enabled", True)]
    names = [source["name"] for source in sources]
    if not names or len(names) != len(set(names)):
        raise ValueError("Configure at least one source, with unique source names")
    previous = read_json(state_path, [])
    review_path = output / "reviews.json"
    reviews = read_json(review_path, {})
    workbook_path = output / "opportunities.xlsx"
    if workbook_path.exists():
        reviews.update({row["id"]: {"decision": row["decision"], "notes": row["notes"]}
                        for row in read_review(workbook_path)})
    batches, statuses = {}, []
    for source in sources:
        try:
            rows = collect_source(source, config, config_root)
            batches[source["name"]] = [row for row in rows if eligible(row, config, today)]
            statuses.append({"name": source["name"], "status": "OK", "count": len(batches[source["name"]])})
        except Exception as error:
            # Keep previous observations if a source fails; expose the failure in outputs and exit status.
            statuses.append({"name": source["name"], "status": "Failed", "error": str(error)})
            print(f'Source failed: {source["name"]}: {error}', file=sys.stderr)
    report = {"collected_on": today.isoformat(), "sources": statuses}
    output.mkdir(parents=True, exist_ok=True)
    write_json(output / "collection-report.json", report)
    if not batches:
        raise ValueError("Every source failed. Existing state and workbook were preserved.")
    records = merge(previous, batches, set(names), today)
    records = [row for row in records if row.get("community_record") or
               reviews.get(row["id"], {}).get("notes", "").strip() or
               (undergraduate_internship(row) and posted_recently(row, today, config.get("max_age_days", 7)))]
    visible = [row for row in records if undergraduate_internship(row)
               and posted_recently(row, today, config.get("max_age_days", 7))]
    if automate:
        selected, screening = screen_records(visible, reviews, config, output, today, refresh=refresh_pages)
        report["screening"] = screening
        count = prepare_digest(selected, output / "digest", config.get("club_name", "Hack@Davidson"), today,
                               config.get("stale_after_days", 7), config.get("max_age_days", 7), allow_empty=True)
        print(f"Prepared {count} opportunities. Email preview: {output / 'digest/email-preview.html'}")
        print(f"Screening details: {output / 'screening-report.html'}")
    export_workbook(visible, reviews, report, workbook_path, config, today)
    export_csv(visible, output / "opportunities.csv")
    write_catalog(records, config_root / "OPPORTUNITIES.md", today)
    write_json(review_path, reviews)
    write_json(state_path, records)
    print(f"Exported {len(visible)} undergraduate internships ({len(records)} current/community records stored). Excel review: {workbook_path}")
    return 1 if any(row["status"] == "Failed" for row in statuses) else 0


def mark_sent(output: Path) -> int:
    """Record only the IDs in the last generated draft, after the user sends it."""
    from openpyxl import load_workbook

    report = read_json(output / "digest/digest-report.json", {})
    ids = set(report.get("included_ids", []))
    if not ids:
        raise ValueError("The last draft contains no listings to mark Sent")
    path = output / "opportunities.xlsx"
    rows = read_review(path)
    reviews = read_json(output / "reviews.json", {})
    reviews.update({row["id"]: {"decision": row["decision"], "notes": row["notes"]} for row in rows})
    for ident in ids:
        reviews[ident] = {**reviews.get(ident, {}), "decision": "Sent"}
    book = load_workbook(path)
    try:
        sheet = book["Opportunities"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        for row in sheet.iter_rows(min_row=2):
            if row[headers["ID"] - 1].value in ids:
                row[headers["Decision"] - 1].value = "Sent"
        temporary = path.with_suffix(".tmp.xlsx")
        book.save(temporary)
        temporary.replace(path)
    finally:
        book.close()
    write_json(output / "reviews.json", reviews)
    print(f"Marked {len(ids)} listings Sent. No email was sent by this command.")
    return 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Collect opportunities, review in Excel, and prepare a club email.")
    sub = parser.add_subparsers(dest="command", required=True)
    gather = sub.add_parser("collect", help="Refresh listings and generate the Excel review workbook")
    gather.add_argument("--config", type=Path, default=Path("config.json"))
    gather.add_argument("--state", type=Path, default=Path("data/opportunities.json"))
    gather.add_argument("--output", type=Path, default=Path("output"))
    run = sub.add_parser("run", help="Collect, screen employer pages, and prepare a ranked email draft")
    run.add_argument("--config", type=Path, default=Path("config.json"))
    run.add_argument("--state", type=Path, default=Path("data/opportunities.json"))
    run.add_argument("--output", type=Path, default=Path("output"))
    run.add_argument("--max-pages", type=int, help="Maximum new employer pages to check this run")
    run.add_argument("--refresh-pages", action="store_true", help="Recheck pages even if cached today")
    sent = sub.add_parser("mark-sent", help="After sending, mark the last draft's listings Sent")
    sent.add_argument("--output", type=Path, default=Path("output"))
    digest = sub.add_parser("digest", help="Build email drafts from Include rows in a saved workbook")
    digest.add_argument("--workbook", type=Path, default=Path("output/opportunities.xlsx"))
    digest.add_argument("--config", type=Path, default=Path("config.json"))
    digest.add_argument("--output", type=Path, default=Path("output/digest"))
    args = parser.parse_args(argv)
    today = datetime.now(timezone.utc).date()
    try:
        if args.command == "mark-sent":
            return mark_sent(args.output)
        config = json.loads(args.config.read_text(encoding="utf-8"))
        if args.command in {"collect", "run"}:
            if args.command == "run" and args.max_pages is not None:
                config.setdefault("screening", {})["max_pages_per_run"] = args.max_pages
            return collect(config, args.config.resolve().parent, args.state, args.output, today,
                           automate=args.command == "run", refresh_pages=getattr(args, "refresh_pages", False))
        count = prepare_digest(read_review(args.workbook), args.output,
                               config.get("club_name", "Hack@Davidson"), today,
                               config.get("stale_after_days", 7),
                               max_age_days=config.get("max_age_days", 7))
        print(f"Prepared {count} opportunities in {args.output / 'email-preview.html'}. No email was sent.")
        return 0
    except (ValueError, OSError, KeyError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
