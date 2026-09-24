from __future__ import annotations

import csv
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .core import availability, clean, iso_date, web_url

FIELDS = [
    ("decision", "Decision", 16), ("notes", "Club notes", 42),
    ("title", "Opportunity", 55), ("organization", "Organization", 28),
    ("category", "Category", 22), ("location", "Location", 36),
    ("deadline", "Deadline", 16), ("availability", "Availability", 23),
    ("url", "Application link", 48), ("eligibility", "Eligibility", 65),
    ("published_date", "Published", 16), ("first_seen", "First seen", 16),
    ("last_seen", "Last seen", 16), ("source_url", "Source link", 48),
    ("sources", "Sources", 40), ("id", "ID", 28),
]
DATES = {"deadline", "published_date", "first_seen", "last_seen"}
DECISIONS = {"Pending", "Include", "Skip", "Sent"}


def read_review(path: Path) -> list[dict]:
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if "Opportunities" not in book.sheetnames:
            raise ValueError("Workbook is missing its Opportunities sheet")
        rows = book["Opportunities"].iter_rows()
        headers = [cell.value for cell in next(rows)]
        expected = [header for _, header, _ in FIELDS]
        if headers != expected:
            raise ValueError("Workbook columns changed. Keep the exported header row intact.")
        result, seen = [], set()
        for number, cells in enumerate(rows, 2):
            if all(cell.value is None for cell in cells):
                continue
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError(f"Row {number}: formulas are not supported in review fields")
            row = {key: clean(cell.value) for (key, _, _), cell in zip(FIELDS, cells)}
            for (key, _, _), cell in zip(FIELDS, cells):
                if key in DATES:
                    row[key] = iso_date(cell.value)
            row["decision"] = row["decision"] or "Pending"
            if row["decision"] not in DECISIONS:
                raise ValueError(f"Row {number}: choose Pending, Include, Skip, or Sent")
            if not row["id"] or row["id"] in seen:
                raise ValueError(f"Row {number}: missing or duplicate ID")
            seen.add(row["id"])
            web_url(row["url"])
            web_url(row["source_url"])
            if not row["title"] or not row["last_seen"]:
                raise ValueError(f"Row {number}: missing title or last-seen date")
            result.append(row)
        return result
    finally:
        book.close()


def literal(cell, value) -> None:
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Never execute source text as an Excel formula.


def export_workbook(records: list[dict], reviews: dict, report: dict, path: Path, config: dict, today: date) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "Opportunities"
    sheet.append([header for _, header, _ in FIELDS])
    for row_index, record in enumerate(records, 2):
        row = {**record, **reviews.get(record["id"], {})}
        row["decision"] = row.get("decision", "Pending")
        row["availability"] = availability(record, today, config.get("stale_after_days", 7))
        row["sources"] = "; ".join(record.get("active_sources", []))
        for column, (key, _, _) in enumerate(FIELDS, 1):
            value = row.get(key, "")
            cell = sheet.cell(row_index, column)
            literal(cell, date.fromisoformat(value) if key in DATES and value else value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if key in DATES:
                cell.number_format = "yyyy-mm-dd"
            if key in {"url", "source_url"} and value:
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
            elif key in {"decision", "notes"}:
                cell.fill = PatternFill("solid", fgColor="EDF5FF")
                cell.font = Font(color="174EA6")
        sheet.row_dimensions[row_index].height = 60
    for column, (_, _, width) in enumerate(FIELDS, 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:P{max(sheet.max_row, 1)}"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    if records:
        table = Table(displayName="OpportunitiesTable", ref=f"A1:P{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    validation = DataValidation(type="list", formula1='"Pending,Include,Skip,Sent"')
    validation.errorTitle = "Choose a review decision"
    validation.error = "Use Pending, Include, Skip, or Sent."
    validation.showErrorMessage = True
    validation.errorStyle = "stop"
    sheet.add_data_validation(validation)
    validation.add(f"A2:A{max(1000, sheet.max_row + 100)}")
    if report.get("screening"):
        screening = book.create_sheet("Screening")
        screening.append(["Selected for email", "Match", "Priority score", "Opportunity", "Organization", "Why", "Unknowns", "Authorization", "Class years", "Skills mentioned", "Pay excerpt", "Checked", "Application link", "Evidence"])
        for result in report["screening"]["results"]:
            screening.append(["Yes" if result["selected"] else "", result["status"], result["score"], result["title"], result["organization"],
                              "; ".join(result["reasons"]), "; ".join(result["unknowns"]), result["authorization"], result["class_years"],
                              ", ".join(result["skills"]), result["pay"], result["checked_on"], result["url"],
                              "\n".join(f"{key}: {value}" for key, value in result["evidence"].items() if value)])
        screening.freeze_panes = "D2"
        screening.auto_filter.ref = screening.dimensions
        for cells in screening:
            screening.row_dimensions[cells[0].row].height = 60
            for cell in cells:
                literal(cell, cell.value)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="17324D")
                    cell.font = Font(color="FFFFFF", bold=True)
        for index in range(1, 15):
            screening.column_dimensions[screening.cell(1, index).column_letter].width = 24 if index < 4 else 50
        for cell in list(screening.columns)[12][1:]:
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")
    guide = book.create_sheet("Instructions")
    instructions = [
        ("Hack@Davidson opportunity review", ""),
        ("Automatic workflow", "Run python -m opportunities run to collect, screen, and generate a compact 20-opportunity email draft. Open output/digest/email-preview.html and use Copy formatted email."),
        ("Screening", "See the Screening tab or output/screening-report.html for reasons, evidence, and unknowns. Priority scores are not acceptance probabilities. Unchecked pages are not automatically selected."),
        ("Manual overrides", "Pending lets automatic screening decide. Include gives a row priority; Skip and Sent exclude it. Include still respects known hard exclusions and the posting window."),
        ("After sending", "Run python -m opportunities mark-sent to mark only the last draft's listings Sent. Neither command sends email."),
        ("1. Review", "Filter Availability to Current. Verify eligibility and deadlines on the application page."),
        ("2. Select", "Set Decision to Include for listings you want in the next email."),
        ("3. Customize", "Add your own wording in Club notes. Keep IDs and column headers intact."),
        ("4. Save", "Save this file as an Excel workbook (.xlsx)."),
        ("5. Create email", "Run: python -m opportunities digest --workbook PATH_TO_YOUR_WORKBOOK"),
        ("6. Send", "Open the email preview, copy the formatted email, add the club recipient, and send from your mail app."),
        ("7. Record", "After sending, set those rows to Sent and save. Generating a draft does not mark anything as sent."),
        ("Refresh", "Local collection preserves Decision and Club notes in the existing output workbook. Downloaded daily workbooks start a new review."),
        ("Blank deadlines", "A blank deadline means the source did not provide one. It does not mean applications remain open indefinitely."),
        ("Availability", "Current means present in a recent source snapshot; application links and eligibility still need review."),
        ("Source coverage", "Only undergraduate internships/co-ops are included. Listings need explicit undergraduate or bachelor's eligibility; graduate/full-time titles and unknown degrees are excluded."),
        ("Collected on", today.isoformat()),
    ]
    for values in instructions:
        guide.append(values)
    health = book.create_sheet("Sources")
    health.append(["Source", "Result", "Accepted listings", "Details"])
    for source in report["sources"]:
        health.append([source["name"], source["status"], source.get("count", 0), source.get("error", "")])
    for tab, widths in [(guide, [30, 115]), (health, [42, 18, 22, 100])]:
        tab.freeze_panes = "A2"
        tab.sheet_view.showGridLines = False
        for column, width in enumerate(widths, 1):
            tab.column_dimensions[tab.cell(1, column).column_letter].width = width
        for cells in tab:
            tab.row_dimensions[cells[0].row].height = 44
            for cell in cells:
                literal(cell, cell.value)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="17324D")
                    cell.font = Font(color="FFFFFF", bold=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    book.save(temporary)
    temporary.replace(path)


def export_csv(records: list[dict], path: Path) -> None:
    keys = ["title", "organization", "category", "location", "deadline", "url", "source_url", "eligibility", "first_seen", "last_seen", "id"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        for row in records:
            # CSV readers may otherwise interpret untrusted values as formulas.
            writer.writerow({key: "'" + str(row.get(key, "")) if str(row.get(key, "")).lstrip().startswith(("=", "+", "-", "@"))
                             else row.get(key, "") for key in keys})
