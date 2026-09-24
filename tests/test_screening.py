import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from opportunities.__main__ import collect, main, mark_sent
from opportunities.core import normalize
from opportunities.digest import prepare_digest
from opportunities.job_pages import JobReader, parse_page, public_url
from opportunities.screening import assess, screen_records
from opportunities.workbook import read_review

TODAY = date(2026, 9, 22)
DESCRIPTION = """Software Engineering Internship
You must be currently pursuing a bachelor's degree in computer science.
Build software using Python, Java and SQL. Work alongside a mentor on data analytics projects.
This paid internship offers $30 per hour. Sophomores and juniors are encouraged to apply.
We cannot provide visa sponsorship for this internship.
"""


def record(ident="one", **overrides):
    return {**normalize({"title": "Software Engineering Intern", "organization": "Example",
                         "category": "Internship", "eligibility": "Bachelor's students",
                         "url": f"https://example.org/{ident}", "published_date": TODAY.isoformat(), **overrides}),
            "first_seen": TODAY.isoformat(), "last_seen": TODAY.isoformat(), "active_sources": ["Test"]}


def page(text=DESCRIPTION, **overrides):
    return {"title": "Software Engineering Intern", "text": text, "posted": TODAY.isoformat(),
            "deadline": "", "format": "Test API", **overrides}


class PageReaderTests(unittest.TestCase):
    def test_extracts_matching_jobposting_and_ignores_scripts(self):
        metadata = {"@context": "https://schema.org", "@type": "JobPosting", "title": "Software Engineering Intern",
                    "description": "<p>Bachelor's students</p><p>Python internship</p>", "datePosted": "2026-09-22"}
        body = '<script>alert("ignore me")</script><script type="application/ld+json">' + json.dumps(metadata) + '</script>'
        result = parse_page(body, "Software Engineering Intern")
        self.assertEqual(result["posted"], "2026-09-22")
        self.assertIn("Bachelor's", result["text"])
        self.assertNotIn("alert", result["text"])

    def test_unrelated_job_and_challenge_pages_are_not_evidence(self):
        with self.assertRaises(ValueError):
            parse_page('<script type="application/ld+json">{"@type":"JobPosting","title":"Accountant"}</script>', "Software Intern")
        with self.assertRaises(ValueError):
            parse_page("<p>Verify you are human</p>", "Software Intern")

    def test_private_addresses_are_rejected(self):
        with patch("opportunities.job_pages.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("127.0.0.1", 443))]):
            with self.assertRaisesRegex(ValueError, "public internet"):
                public_url("https://example.org/job")

    def test_greenhouse_adapter_uses_description_and_original_publication(self):
        row = record(url="https://job-boards.greenhouse.io/example/jobs/123")
        response = {"title": row["title"], "content": '<p>' + DESCRIPTION + '</p>',
                    "first_published": "2026-09-22T00:00:00Z", "updated_at": "2026-09-23"}
        with patch.object(JobReader, "read", return_value=json.dumps(response)) as read:
            actual = JobReader().job(row)
        self.assertEqual(actual["posted"], response["first_published"])
        read.assert_called_once_with("https://boards-api.greenhouse.io/v1/boards/example/jobs/123", api=True)

    def test_lever_includes_requirements_and_ashby_matches_job_id(self):
        row = record(url="https://jobs.lever.co/example/abc")
        response = {"text": row["title"], "descriptionPlain": "Internship", "lists": [{"text": "Requirements", "content": DESCRIPTION}]}
        with patch.object(JobReader, "read", return_value=json.dumps(response)):
            self.assertIn("bachelor", JobReader().job(row)["text"])
        row = record(url="https://jobs.ashbyhq.com/example/abc/application")
        response = {"jobs": [{"title": row["title"], "descriptionPlain": DESCRIPTION, "jobUrl": "https://jobs.ashbyhq.com/example/abc", "isListed": True, "publishedAt": "2026-09-22"}]}
        with patch.object(JobReader, "read", return_value=json.dumps(response)):
            self.assertEqual(JobReader().job(row)["posted"], "2026-09-22")


class ScreeningTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_strong_match_has_evidence_and_does_not_infer_authorization(self):
        result = assess(record(), page(), {}, TODAY)
        self.assertEqual(result["status"], "Strong match")
        self.assertIn("bachelor", result["evidence"]["degree"])
        self.assertIn("no sponsorship", result["authorization"])
        self.assertNotIn("eligible", result["authorization"])
        self.assertIn("python", result["skills"])
        self.assertIn("Sophomores", result["class_years"])

    def test_davidson_fit_and_location_receive_priority(self):
        ordinary = assess(record(organization="Ordinary Co", location="Boston, MA"), page(), {}, TODAY)
        nearby = assess(record(organization="Ordinary Co", location="Charlotte, NC"), page(), {}, TODAY)
        state = assess(record(organization="Ordinary Co", location="Raleigh, NC"), page(), {}, TODAY)
        self.assertEqual(nearby["score"], ordinary["score"] + 20)
        self.assertEqual(state["score"], ordinary["score"] + 12)
        self.assertIn("Davidson student fit", ordinary["priority_labels"])
        self.assertIn("Class-year fit", ordinary["priority_labels"])

    def test_email_screening_uses_readme_davidson_priority_signals(self):
        config = {"readme_priority": {
            "north_carolina_bonus": 120,
            "nearby_states": {"sc": 80},
            "davidson_employers": ["Trane Technologies"],
            "davidson_employer_bonus": 160,
        }}
        ordinary = assess(record(organization="Ordinary Co", location="Seattle, WA"), page(), config, TODAY)
        nearby = assess(record(organization="Ordinary Co", location="Charlotte, NC"), page(), config, TODAY)
        connected = assess(record(organization="Trane Technologies", location="La Crosse, WI"), page(), config, TODAY)
        self.assertGreater(nearby["score"], ordinary["score"])
        self.assertGreater(connected["score"], nearby["score"])
        self.assertIn("Davidson-connected employer", connected["priority_labels"])

    def test_email_order_uses_readme_priority_before_page_fit_score(self):
        local = record("local", organization="Local NC", location="Charlotte, NC")
        distant = record("distant", organization="Distant Co", location="Seattle, WA")
        config = {"readme_priority": {"north_carolina_bonus": 120},
                  "screening": {"max_pages_per_run": 2, "max_email_items": 2}}
        with patch.object(JobReader, "job", side_effect=[page(), page()]):
            selected, _ = screen_records([distant, local], {}, config, self.root, TODAY)
        self.assertEqual(selected[0]["id"], local["id"])

    def test_priority_roles_are_checked_first_with_a_small_budget(self):
        ordinary = record("ordinary", organization="Ordinary Co", location="Boston, MA")
        priority = record("priority", organization="Local Co", location="Charlotte, NC")
        with patch.object(JobReader, "job", return_value=page()) as fetch:
            screen_records([ordinary, priority], {}, {"screening": {"max_pages_per_run": 1}}, self.root, TODAY)
        self.assertEqual(fetch.call_args.args[0]["id"], priority["id"])

    def test_unreadable_and_unknown_degree_pages_never_become_strong(self):
        self.assertEqual(assess(record(), {}, {}, TODAY)["status"], "Not checked")
        self.assertEqual(assess(record(), {"error": "HTTP 403"}, {}, TODAY)["status"], "Needs review")
        self.assertEqual(assess(record(), page("Python software internship; see application for details"), {}, TODAY)["status"], "Needs review")

    def test_employer_age_degree_and_closure_override_feed(self):
        self.assertEqual(assess(record(), page(posted="2026-08-01"), {}, TODAY)["status"], "Excluded")
        self.assertEqual(assess(record(), page(deadline="2026-09-21"), {}, TODAY)["status"], "Excluded")
        self.assertEqual(assess(record(), page("You must be pursuing a PhD in computer science. Software internship."), {}, TODAY)["status"], "Excluded")
        self.assertEqual(assess(record(), page("This job is closed."), {}, TODAY)["status"], "Excluded")

    def test_budget_cache_and_manual_decisions(self):
        records = [record(str(index), title=f"Software Engineering Intern {index}") for index in range(5)]
        reviews = {records[0]["id"]: {"decision": "Sent"}, records[1]["id"]: {"decision": "Skip"},
                   records[2]["id"]: {"decision": "Include", "notes": "My note"}}
        config = {"screening": {"max_pages_per_run": 1, "max_email_items": 2, "max_per_organization": 2}}
        with patch.object(JobReader, "job", return_value=page()) as fetch:
            selected, report = screen_records(records, reviews, config, self.root, TODAY)
            self.assertEqual(fetch.call_count, 1)
            self.assertEqual([r["id"] for r in selected], [records[2]["id"]])
            self.assertEqual(selected[0]["notes"], "My note")
            selected, report = screen_records(records, reviews, config, self.root, TODAY)
            self.assertEqual(fetch.call_count, 2)  # New page; the first page came from cache.
            self.assertEqual(len(selected), 2)
            self.assertEqual(report["cached"], 1)
        self.assertNotIn(records[0]["id"], [row["id"] for row in selected])
        self.assertNotIn(records[1]["id"], [row["id"] for row in selected])

    def test_manual_include_still_requires_verified_strong_match(self):
        row = record()
        reviews = {row["id"]: {"decision": "Include"}}
        with patch.object(JobReader, "job", return_value={}):
            selected, _ = screen_records([row], reviews, {"screening": {"max_pages_per_run": 1}}, self.root, TODAY)
        self.assertEqual(selected, [])

    def test_selected_row_carries_verified_employer_deadline(self):
        row = record()
        with patch.object(JobReader, "job", return_value=page(deadline="2026-09-30")):
            selected, _ = screen_records([row], {}, {"screening": {"max_pages_per_run": 1}}, self.root, TODAY)
        self.assertEqual(selected[0]["deadline"], "2026-09-30")

    def test_employer_diversity_and_scoring_changes_reuse_page_cache(self):
        records = [record(str(index), organization="First" if index < 3 else "Second") for index in range(4)]
        config = {"screening": {"max_per_organization": 1}}
        with patch.object(JobReader, "job", return_value=page()) as fetch:
            selected, report = screen_records(records, {}, config, self.root, TODAY)
            self.assertEqual(len(selected), 2)
            config["screening"]["interests"] = ["marine biology"]
            selected, _ = screen_records(records, {}, config, self.root, TODAY)
            self.assertEqual(selected, [])
            self.assertEqual(fetch.call_count, 4)

    def test_selects_exactly_twenty_when_enough_verified_matches_exist(self):
        records = [record(str(index), title=f"Software Engineering Intern {index}", organization=f"Company {index}")
                   for index in range(25)]
        config = {"screening": {"max_email_items": 20, "max_pages_per_run": 25}}
        with patch.object(JobReader, "job", return_value=page()):
            selected, report = screen_records(records, {}, config, self.root, TODAY)
        self.assertEqual(len(selected), 20)
        self.assertEqual(report["target"], 20)
        self.assertEqual(report["shortfall"], 0)

    def test_empty_draft_clears_prior_artifacts(self):
        output = self.root / "digest"
        row = {**record(), "decision": "Include", "notes": "", "availability": "Current"}
        prepare_digest([row], output, "Club", TODAY)
        self.assertTrue((output / "email-draft.eml").exists())
        self.assertEqual(prepare_digest([], output, "Club", TODAY, allow_empty=True), 0)
        self.assertFalse((output / "email-draft.eml").exists())
        self.assertEqual(json.loads((output / "digest-report.json").read_text())["included_ids"], [])

    def test_end_to_end_automatic_draft_then_mark_sent(self):
        row = record()
        config = {"sources": [{"name": "Test"}]}
        state = self.root / "state.json"
        with patch("opportunities.__main__.collect_source", return_value=[row]), patch.object(JobReader, "job", return_value=page()):
            self.assertEqual(collect(config, self.root, state, self.root, TODAY, automate=True), 0)
        draft = self.root / "digest/email-preview.html"
        self.assertIn("Copy formatted email", draft.read_text())
        self.assertIn(">Apply</a>", draft.read_text())
        self.assertNotIn("Why selected", draft.read_text())
        self.assertEqual(read_review(self.root / "opportunities.xlsx")[0]["decision"], "Pending")
        book = load_workbook(self.root / "opportunities.xlsx")
        self.assertIn("Screening", book.sheetnames)
        book.close()
        mark_sent(self.root)
        self.assertEqual(read_review(self.root / "opportunities.xlsx")[0]["decision"], "Sent")
        with patch("opportunities.__main__.collect_source", return_value=[row]), patch.object(JobReader, "job") as fetch:
            collect(config, self.root, state, self.root, TODAY, automate=True)
            fetch.assert_not_called()
        self.assertFalse((self.root / "digest/email-draft.eml").exists())

    def test_html_report_escapes_untrusted_evidence(self):
        with patch.object(JobReader, "job", return_value=page(DESCRIPTION + "\n<script>alert(1)</script>")):
            screen_records([record(organization="<script>alert(1)</script>")], {}, {}, self.root, TODAY)
        report = (self.root / "screening-report.html").read_text()
        self.assertNotIn("<script>", report)
        self.assertIn("&lt;script&gt;", report)


if __name__ == "__main__":
    unittest.main()
